# Copyright (c) 2026 Edwin Njeru. Licensed under the MIT License (see LICENSE).

"""
Shared openpyxl drawing + formula helpers for bizplan financial renderers.

Ported from the standalone `scripts/xl_helpers.py` (proven, already used by the original
Grand Ballroom Nairobi scripts) so the `bizplan` package's renderers can emit live Excel
formulas instead of Python-computed static values.

FMI data-provenance colour convention:
    BLUE   — input / hardcoded fact (company disclosure)
    DARK   — internal formula
    TEAL   — cross-sheet reference
    ORANGE — modeled proxy / benchmark assumption, not directly disclosed
"""

from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Calibri"

NAVY       = "1F3864"
MED_BLUE   = "2F5597"
LIGHT_BLUE = "BDD7EE"
V_LIGHT    = "DEEAF1"
WHITE      = "FFFFFF"
DARK       = "1F1F1F"
MID_GRAY   = "595959"
LIGHT_GRAY = "F2F2F2"
ALT_ROW    = "EBF3FB"
GOLD       = "C9A227"
TOTAL_FILL = "D6E4F7"
GREEN_DRK  = "375623"
GREEN_LT   = "E2EFDA"
RED_DARK   = "C00000"
ORANGE     = "ED7D31"
TEAL       = "00695C"
BLUE_INPUT = "1F4E96"  # FMI convention: font colour for disclosed/hardcoded input cells


# ── Low-level primitives ───────────────────────────────────────────────────

def fill(hex_color):
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

def font(bold=False, size=10, color=DARK, italic=False, name=FONT_NAME):
    return Font(name=name, bold=bold, size=size, color=color, italic=italic)

def align(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def thin_border(bottom=False, top=False):
    side = Side(style="thin", color="8496AF")
    return Border(
        bottom=side if bottom else Side(style=None),
        top=side    if top    else Side(style=None),
    )

def top_bottom_border():
    side = Side(style="medium", color=NAVY)
    return Border(top=side, bottom=side)

def bottom_only():
    return Border(bottom=Side(style="thin", color="8496AF"))

def box_border(color=NAVY):
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)

def outline_range(ws, row, cols, color=NAVY):
    """Draws a single bounding rectangle around a contiguous horizontal group of cells —
    left edge only on the first column, right edge only on the last, top+bottom on every
    column — instead of `box_border()`'s every-cell-gets-all-four-sides, which reads as a
    grid of separate boxes rather than one outlined group."""
    side = Side(style="thin", color=color)
    none = Side(style=None)
    for i, col in enumerate(cols):
        cell = ws.cell(row=row, column=col)
        cell.border = Border(
            left=side if i == 0 else none,
            right=side if i == len(cols) - 1 else none,
            top=side,
            bottom=side,
        )

def set_col_widths(ws, widths):
    for col_letter, w in widths.items():
        ws.column_dimensions[col_letter].width = w

def set_row_height(ws, row, height):
    ws.row_dimensions[row].height = height


# ── Cell writer ────────────────────────────────────────────────────────────

def write(ws, row, col, value, bold=False, size=10, txt_color=DARK,
          bg=None, halign="left", valign="center", italic=False,
          num_fmt=None, wrap=False, border=None):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    cell.font = font(bold=bold, size=size, color=txt_color, italic=italic)
    cell.alignment = align(h=halign, v=valign, wrap=wrap)
    if bg:
        cell.fill = fill(bg)
    if num_fmt:
        cell.number_format = num_fmt
    if border:
        cell.border = border
    return cell

def num(ws, row, col, value, fmt='#,##0.0', bold=False, bg=None,
        txt_color=DARK, border=None):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    cell.font = font(bold=bold, size=10, color=txt_color)
    cell.alignment = align(h="right", v="center")
    cell.number_format = fmt
    if bg:
        cell.fill = fill(bg)
    if border:
        cell.border = border
    return cell

def pct(ws, row, col, value, bold=False, bg=None, txt_color=DARK):
    return num(ws, row, col, value, fmt='0.0%', bold=bold, bg=bg, txt_color=txt_color)


# ── Row-level builders ─────────────────────────────────────────────────────

def header_row(ws, row, label, sub=None, height=18, bg=NAVY,
               txt_color=WHITE, merge_to_col=None, start_col=2):
    ws.row_dimensions[row].height = height
    for c in range(1, 17):
        ws.cell(row=row, column=c).fill = fill(bg)
    write(ws, row, start_col, label, bold=True, size=11,
          txt_color=txt_color, bg=bg, halign="left")
    if merge_to_col:
        ws.merge_cells(start_row=row, start_column=start_col,
                       end_row=row, end_column=merge_to_col)

def section_header(ws, row, label, bg=MED_BLUE, txt_color=WHITE, height=15):
    ws.row_dimensions[row].height = height
    for c in range(1, 17):
        ws.cell(row=row, column=c).fill = fill(bg)
    write(ws, row, 3, label, bold=True, size=10, txt_color=txt_color, bg=bg)

def year_header_row(ws, row, years, data_cols, bg=MED_BLUE, txt_color=WHITE,
                    projected_label="Projected", label_col=None, units_col=None):
    ws.row_dimensions[row].height = 14
    for c in range(1, 17):
        ws.cell(row=row, column=c).fill = fill(bg)
    if label_col:
        write(ws, row, label_col, "", bg=bg)
    for i, yr in enumerate(years):
        write(ws, row, data_cols[i], str(yr), bold=True, size=10,
              txt_color=txt_color, bg=bg, halign="center")

def blank_row(ws, row, height=5, bg=WHITE):
    ws.row_dimensions[row].height = height

def alt_fill(row_index):
    return ALT_ROW if row_index % 2 == 0 else WHITE

def data_row(ws, row, label, values, data_cols, indent=0,
             bold=False, bg=None, txt_color=DARK, fmt='#,##0.0',
             units=None, units_col=None, label_col=3, alt_idx=None,
             border=None, italic=False):
    ws.row_dimensions[row].height = 14
    row_bg = bg if bg else (alt_fill(alt_idx) if alt_idx is not None else WHITE)
    for c in range(1, 17):
        ws.cell(row=row, column=c).fill = fill(row_bg)
    lbl = ("    " * indent) + label if indent else label
    write(ws, row, label_col, lbl, bold=bold, txt_color=txt_color,
          bg=row_bg, italic=italic)
    if units and units_col:
        write(ws, row, units_col, units, bold=False, txt_color=MID_GRAY,
              bg=row_bg, halign="center", italic=True, size=9)
    for i, v in enumerate(values):
        if v is not None:
            num(ws, row, data_cols[i], v, fmt=fmt, bold=bold,
                bg=row_bg, txt_color=txt_color, border=border)

def total_row(ws, row, label, values, data_cols, bg=TOTAL_FILL,
              txt_color=DARK, fmt='#,##0.0', label_col=3, border=True):
    ws.row_dimensions[row].height = 15
    for c in range(1, 17):
        ws.cell(row=row, column=c).fill = fill(bg)
    write(ws, row, label_col, label, bold=True, txt_color=txt_color, bg=bg)
    b = top_bottom_border() if border else None
    for i, v in enumerate(values):
        if v is not None:
            num(ws, row, data_cols[i], v, fmt=fmt, bold=True,
                bg=bg, txt_color=txt_color, border=b)


# ── Formula helpers ────────────────────────────────────────────────────────
# Uses get_column_letter (handles columns beyond Z, unlike chr(64+col)).

def _cell(row, col):
    return f"{get_column_letter(col)}{row}"

def _sum_f(r_start, r_end, col):
    return f"=SUM({_cell(r_start, col)}:{_cell(r_end, col)})"

def _add_rows_f(rows, col):
    return "=" + "+".join(_cell(r, col) for r in rows)

def _sub_f(r_a, r_b, col):
    return f"={_cell(r_a, col)}-{_cell(r_b, col)}"

def _ratio_f(r_num, r_denom, col):
    return f"=IF({_cell(r_denom, col)}<>0,{_cell(r_num, col)}/{_cell(r_denom, col)},0)"

def _ref_f(row, col):
    return f"={_cell(row, col)}"

def _model_ref(row, col):
    return f"='Model'!{_cell(row, col)}"
